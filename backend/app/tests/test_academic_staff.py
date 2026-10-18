import io
import openpyxl
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework import status
from django.core.files.uploadedfile import SimpleUploadedFile

from app.models import (
    CustomUser,
    AcademicStaff,
    AcademicBatch,
    CourseClass,
    Student,
    Supervisor,
    GraduationProject,
    FinalGradeSummary,
    AuditLog,
)
from app.services.excel_importer import ExcelImportService


class AcademicStaffModuleTestCase(TestCase):
    def setUp(self):
        self.client = APIClient()

        # 1. Create Academic Staff user & profile
        self.academic_staff_user = CustomUser.objects.create_user(
            username="test_giao_vu",
            email="giaovu_test@utc.edu.vn",
            password="GiaovuPassword123!",
            user_type="academic_staff",
            is_staff=True,
            first_name="Văn",
            last_name="Giáo Vụ",
        )
        self.academic_staff_profile = AcademicStaff.objects.create(
            user=self.academic_staff_user,
            staff_id="GVU_TEST01",
            phone_number="0987654321",
            office_location="Văn phòng Khoa CNTT (P.405-A9)",
        )

        # 2. Create Admin user for comparison/admin actions
        self.admin_user = CustomUser.objects.create_user(
            username="test_admin",
            email="admin_test@utc.edu.vn",
            password="AdminPassword123!",
            user_type="admin",
            is_staff=True,
        )

        # 3. Create Batch & CourseClass
        self.batch = AcademicBatch.objects.create(
            batch_code="2026_K62_GVU",
            batch_name="Đợt ĐATN K62 Khoa CNTT",
            program_type="BOTH",
            current_stage="PHASE_4_ELIGIBILITY",
            is_active=True,
        )

        self.course_class = CourseClass.objects.create(
            batch=self.batch,
            class_code="CNTT1-K62",
            class_name="Công nghệ thông tin 1 K62",
            program_type="DAI_TRA",
            education_program="BACHELOR",
        )

        # 4. Create Supervisor
        self.sup_user = CustomUser.objects.create_user(
            username="gv_huongdan",
            email="gvhd@utc.edu.vn",
            password="Password123!",
            user_type="supervisor",
            first_name="Hùng",
            last_name="Trần",
        )
        self.supervisor = Supervisor.objects.create(
            user=self.sup_user,
            supervisor_id="GV001",
            academic_title="TS.",
        )

        # 5. Create Student 1: debt_credits = 6 (< 10 -> Eligible)
        self.std_user_1 = CustomUser.objects.create_user(
            username="sv_eligible",
            email="sv1@utc.edu.vn",
            password="StudentPass123!",
            user_type="student",
            first_name="An",
            last_name="Nguyễn Văn",
        )
        self.student_1 = Student.objects.create(
            user=self.std_user_1,
            registration_no="221200001",
            course_class=self.course_class,
            academic_batch=self.batch,
            education_program="BACHELOR",
            department="CNTT",
        )
        self.project_1 = GraduationProject.objects.create(
            student=self.student_1,
            supervisor=self.supervisor,
            batch=self.batch,
            topic_title_vi="Nghiên cứu AI trong quản lý giáo dục",
            debt_credits=6,
            status="OUTLINE_APPROVED",
            initial_eligibility="PENDING",
        )

        # 6. Create Student 2: debt_credits = 14 (>= 10 -> Ineligible)
        self.std_user_2 = CustomUser.objects.create_user(
            username="sv_ineligible",
            email="sv2@utc.edu.vn",
            password="StudentPass123!",
            user_type="student",
            first_name="Bình",
            last_name="Lê Văn",
        )
        self.student_2 = Student.objects.create(
            user=self.std_user_2,
            registration_no="221200002",
            course_class=self.course_class,
            academic_batch=self.batch,
            education_program="BACHELOR",
            department="CNTT",
        )
        self.project_2 = GraduationProject.objects.create(
            student=self.student_2,
            supervisor=self.supervisor,
            batch=self.batch,
            topic_title_vi="Xây dựng sàn thương mại điện tử Microservices",
            debt_credits=14,
            status="OUTLINE_APPROVED",
            initial_eligibility="PENDING",
        )

    # =========================================================================
    # 1. Authentication & Role Tests
    # =========================================================================
    def test_academic_staff_login_successful(self):
        """Giáo vụ đăng nhập thành công vào admin portal với token và profile đúng"""
        response = self.client.post("/app/admin/login/", {
            "email": "giaovu_test@utc.edu.vn",
            "password": "GiaovuPassword123!",
        }, format="json")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("access", response.data)
        self.assertEqual(response.data["user_type"], "academic_staff")
        self.assertIn("academic_staff_profile", response.data["user"])
        profile_data = response.data["user"]["academic_staff_profile"]
        self.assertEqual(profile_data["staff_id"], "GVU_TEST01")
        self.assertEqual(profile_data["phone_number"], "0987654321")

    def test_academic_staff_token_refresh(self):
        """Token refresh giữ nguyên vai trò academic_staff"""
        login_res = self.client.post("/app/admin/login/", {
            "email": "giaovu_test@utc.edu.vn",
            "password": "GiaovuPassword123!",
        }, format="json")
        self.assertEqual(login_res.status_code, status.HTTP_200_OK)

        refresh_res = self.client.post("/app/token/refresh/", {}, format="json")
        self.assertEqual(refresh_res.status_code, status.HTTP_200_OK)
        self.assertIn("access", refresh_res.data)
        self.assertEqual(refresh_res.data["user_type"], "academic_staff")

    # =========================================================================
    # 2. Debt Credit Evaluation Rule (< 10 vs >= 10)
    # =========================================================================
    def test_import_eligibility_by_debt_credits_threshold(self):
        """
        Kiểm tra quy tắc xét điều kiện nợ tín chỉ:
        - Nợ < 10 tín chỉ: Đủ điều kiện (ELIGIBLE), chuyển sang IN_PROGRESS để giao task.
        - Nợ >= 10 tín chỉ: Không đủ điều kiện (INELIGIBLE), giữ OUTLINE_APPROVED chờ Khoa xét duyệt.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "XetDieuKien"

        # Headers
        ws.append(["STT", "Mã SV", "Họ và tên", "Lớp", "Số TC nợ", "Điểm GPA"])
        ws.append([1, "221200001", "Nguyễn Văn An", "CNTT1-K62", 6, 3.2])
        ws.append([2, "221200002", "Lê Văn Bình", "CNTT1-K62", 14, 2.4])

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)

        uploaded_file = SimpleUploadedFile(
            "bang_diem_xet_dieu_kien.xlsx",
            excel_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        res = ExcelImportService.import_eligibility_from_excel(uploaded_file, self.batch.id)
        self.assertTrue(res["success"])
        self.assertEqual(res["total_processed"], 2)
        self.assertEqual(res["eligible_count"], 1)
        self.assertEqual(res["ineligible_count"], 1)

        # Refresh database state
        self.project_1.refresh_from_db()
        self.project_2.refresh_from_db()

        # Kiểm tra SV 1 (Nợ 6 < 10)
        self.assertEqual(self.project_1.debt_credits, 6)
        self.assertEqual(self.project_1.initial_eligibility, "ELIGIBLE")
        self.assertEqual(self.project_1.status, "IN_PROGRESS", "Nợ < 10 TC phải chuyển sang IN_PROGRESS để giao task")

        # Kiểm tra SV 2 (Nợ 14 >= 10)
        self.assertEqual(self.project_2.debt_credits, 14)
        self.assertEqual(self.project_2.initial_eligibility, "INELIGIBLE")
        self.assertEqual(self.project_2.status, "OUTLINE_APPROVED", "Nợ >= 10 TC phải giữ trạng thái chờ xét Force Approve")

    def test_force_approve_and_disqualify_ineligible_project(self):
        """Khoa phê duyệt Force Approve hoặc Loại khỏi đợt với sinh viên nợ >= 10 TC"""
        self.client.force_authenticate(user=self.academic_staff_user)

        # 1. Force approve SV 2
        force_res = self.client.post(
            f"/app/admin/projects/{self.project_2.id}/force-approve/",
            {"decision": "FORCE_APPROVED", "reason": "Khoa đặc cách do đang học trả nợ song song kỳ cuối"},
            format="json"
        )
        self.assertEqual(force_res.status_code, status.HTTP_200_OK)
        self.project_2.refresh_from_db()
        self.assertEqual(self.project_2.initial_eligibility, "FORCE_APPROVED")
        self.assertEqual(self.project_2.status, "IN_PROGRESS")

        # 2. Disqualify (loại khỏi đợt)
        disqualify_res = self.client.post(
            f"/app/admin/projects/{self.project_2.id}/force-approve/",
            {"decision": "DISQUALIFIED", "reason": "Không đáp ứng chuẩn đầu ra và nợ quá số tín chỉ"},
            format="json"
        )
        self.assertEqual(disqualify_res.status_code, status.HTTP_200_OK)
        self.project_2.refresh_from_db()
        self.assertEqual(self.project_2.initial_eligibility, "DISQUALIFIED")
        self.assertEqual(self.project_2.status, "DISQUALIFIED")

    # =========================================================================
    # 3. End-of-Term Features for Academic Staff (Final Grades & Class Export)
    # =========================================================================
    def test_academic_staff_final_grades_aggregation(self):
        """Giáo vụ xem bảng điểm tổng hợp cuối kỳ lọc theo lớp và toàn đợt"""
        self.client.force_authenticate(user=self.academic_staff_user)

        # Setup điểm cho SV 1 (Đạt)
        self.project_1.supervisor_score = 8.5
        self.project_1.reviewer_score = 8.0
        self.project_1.save()
        summary1, _ = FinalGradeSummary.objects.get_or_create(project=self.project_1)
        summary1.supervisor_score = 8.5
        summary1.reviewer_score = 8.0
        summary1.council_avg_score = 8.5
        summary1.calculate_and_save()

        # Gọi API tổng hợp điểm
        response = self.client.get(f"/app/admin/academic-staff/final-grades/?batch_id={self.batch.id}")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        data = response.data
        self.assertEqual(data["batch"]["id"], self.batch.id)
        self.assertEqual(len(data["classes"]), 1)
        self.assertEqual(data["classes"][0]["class_code"], "CNTT1-K62")
        self.assertEqual(len(data["grades"]), 2)

        # Kiểm tra thống kê
        stats = data["stats"]
        self.assertEqual(stats["total_students"], 2)
        self.assertEqual(stats["passed_count"], 1)

    def test_academic_staff_export_class_grades_excel(self):
        """Giáo vụ xuất file Excel bảng điểm theo lớp chuẩn định dạng UTC"""
        self.client.force_authenticate(user=self.academic_staff_user)

        # 1. Xuất theo toàn đợt
        res_all = self.client.get(f"/app/admin/academic-staff/export-class-grades-excel/?batch_id={self.batch.id}")
        self.assertEqual(res_all.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res_all["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("Bang_Diem_Do_An", res_all["Content-Disposition"])

        # Đọc nội dung file Excel trả về bằng openpyxl để xác nhận cấu trúc
        content_stream = io.BytesIO(res_all.content)
        wb = openpyxl.load_workbook(content_stream)
        self.assertIn("Toàn đợt", wb.sheetnames)
        ws = wb["Toàn đợt"]

        # Kiểm tra có tiêu đề và thông tin sinh viên
        cell_values = [cell for row in ws.iter_rows(values_only=True) for cell in row if cell is not None]
        self.assertTrue(any("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI" in str(v) for v in cell_values))
        self.assertTrue(any("221200001" in str(v) for v in cell_values))

        # 2. Xuất theo lớp cụ thể
        res_class = self.client.get(
            f"/app/admin/academic-staff/export-class-grades-excel/?batch_id={self.batch.id}&class_id={self.course_class.id}"
        )
        self.assertEqual(res_class.status_code, status.HTTP_200_OK)
        self.assertIn(f"_Lop_{self.course_class.id}", res_class["Content-Disposition"])

    def test_academic_staff_complete_batch(self):
        """Giáo vụ hoàn tất và khóa sổ đợt đồ án tốt nghiệp"""
        self.client.force_authenticate(user=self.academic_staff_user)

        res = self.client.post("/app/admin/academic-staff/complete-batch/", {
            "batch_id": self.batch.id,
        }, format="json")

        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.batch.refresh_from_db()
        self.assertEqual(self.batch.current_stage, "COMPLETED")
        self.assertTrue(self.batch.is_closed)

        # Kiểm tra audit log
        audit = AuditLog.objects.filter(user=self.academic_staff_user).first()
        self.assertIsNotNone(audit)
        self.assertIn("đóng đợt", audit.description.lower())
