import io
import sys
import openpyxl

if sys.platform.startswith("win"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

wb = openpyxl.load_workbook(r"c:\Users\Admin\source\repos\Smart-Fyp-Admin\Test_Report_ThesisHub.xlsx")
sheet = wb["Test Report"]

print(f"Total rows: {sheet.max_row}, Total cols: {sheet.max_column}")

headers = [cell.value for cell in sheet[1]]
print("Headers:", headers)

for r in range(2, sheet.max_row + 1):
    row_vals = [cell.value for cell in sheet[r]]
    print(f"Row {r}: {row_vals}")
